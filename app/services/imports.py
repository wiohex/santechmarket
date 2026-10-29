import csv
import io
import secrets
import zipfile

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select

from app.models import Category, Product
from app.schemas.forms import ProductForm

REQUIRED = {"sku", "name", "category", "price"}
ALLOWED = REQUIRED | {"in_stock", "description"}


def parse_price_file(filename: str, content: bytes) -> list[dict]:
    if len(content) > 2 * 1024 * 1024:
        raise ValueError("Прайс должен быть не больше 2 МБ.")
    if filename.lower().endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1251")
        delimiter = ";" if ";" in text.splitlines()[0] else ","
        try:
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
        except csv.Error:
            raise ValueError("Некорректный CSV или слишком длинная ячейка.") from None
    elif filename.lower().endswith(".xlsx"):
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if (
                len(archive.infolist()) > 1000
                or sum(i.file_size for i in archive.infolist()) > 20 * 1024 * 1024
            ):
                raise ValueError("Распакованный XLSX слишком большой.")
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = book.worksheets[0]
            if sheet.max_column and sheet.max_column > 20:
                raise ValueError("В таблице слишком много столбцов.")
            # Never silently truncate input, even when worksheet dimensions are missing.
            if sheet.max_row and sheet.max_row > 1001:
                raise ValueError(
                    "В XLSX допускается до 1001 строки, включая заголовок и пустые строки."
                )
            rows = []
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= 1001 or len(row) > 20:
                    raise ValueError("В XLSX слишком много строк или столбцов.")
                rows.append(row)
        finally:
            book.close()
    else:
        raise ValueError("Поддерживаются только CSV и XLSX.")
    rows = [list(row) for row in rows if any(v is not None and str(v).strip() for v in row)]
    if len(rows) < 2:
        raise ValueError("В файле нет строк с товарами.")
    if len(rows) > 1001:
        raise ValueError("За один импорт можно загрузить до 1000 товаров.")
    headers = [str(h or "").strip().lower() for h in rows[0]]
    if (
        len(set(headers)) != len(headers)
        or not REQUIRED.issubset(headers)
        or not set(headers).issubset(ALLOWED)
    ):
        raise ValueError(
            "Столбцы: sku, name, category, price; необязательно in_stock, description. Без повторов и лишних столбцов."
        )
    result, seen = [], set()
    for number, cells in enumerate(rows[1:], start=2):
        if len(cells) != len(headers):
            raise ValueError(f"Строка {number}: число ячеек не совпадает с заголовками.")
        values = {
            key: str(value).strip() if value is not None else ""
            for key, value in zip(headers, cells)
        }
        if any(value.startswith("=") for value in values.values()):
            raise ValueError(f"Строка {number}: замените формулы готовыми значениями.")
        stock = values.get("in_stock", "true").lower()
        if stock not in {"true", "false", "1", "0", "да", "нет", "yes", "no"}:
            raise ValueError(f"Строка {number}: in_stock должен быть 1/0 или да/нет.")
        try:
            item = ProductForm(
                name=values["name"],
                sku=values["sku"],
                category_id=1,
                price=values["price"].replace(",", "."),
                in_stock=stock in {"true", "1", "да", "yes"},
                active=True,
                description=values.get("description", ""),
            )
        except ValidationError:
            raise ValueError(
                f"Строка {number}: проверьте название, артикул и неотрицательную цену (до 2 знаков после запятой)."
            ) from None
        category = values["category"]
        if not 2 <= len(category) <= 100 or item.sku in seen:
            raise ValueError(f"Строка {number}: неверная категория или повтор артикула {item.sku}.")
        seen.add(item.sku)
        result.append(
            {
                "name": item.name,
                "sku": item.sku,
                "category": category,
                "price": int(item.price * 100),
                "in_stock": item.in_stock,
                "description": item.description,
            }
        )
    return result


def apply_import(db, rows: list[dict]):
    categories = {c.name: c for c in db.scalars(select(Category))}
    products = {p.sku: p for p in db.scalars(select(Product))}
    for row in rows:
        if row["category"] not in categories:
            category = Category(name=row["category"], slug="category-" + secrets.token_hex(6))
            db.add(category)
            db.flush()
            categories[category.name] = category
        product = products.get(row["sku"])
        if product is None:
            product = Product(sku=row["sku"])
            db.add(product)
        product.category_id = categories[row["category"]].id
        for key in ("name", "price", "in_stock", "description"):
            setattr(product, key, row[key])
        product.active = True
    db.commit()
