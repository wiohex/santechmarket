from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from conftest import csrf
from openpyxl import Workbook
from sqlalchemy import func, select
from test_admin import login_admin


def test_seed_is_idempotent_and_admin_is_not_overwritten(store):
    from app.models import Product, User
    from scripts.create_admin import create_admin
    from scripts.seed import seed

    client, factory = store
    assert seed(factory) == 12
    assert seed(factory) == 0
    create_admin("new-owner@example.com", "valid-admin-password", factory)
    import pytest

    with pytest.raises(ValueError):
        create_admin("new-owner@example.com", "another-password", factory)
    with factory() as db:
        assert db.scalar(select(User).where(User.email == "new-owner@example.com")).role == "admin"
        assert db.scalar(select(func.count(Product.id))) == 14


def test_corrupted_xlsx_xml_is_validation_error(store):
    client, factory = store
    login_admin(client, factory)
    book = Workbook()
    original = BytesIO()
    book.save(original)
    broken = BytesIO()
    with ZipFile(original) as archive, ZipFile(broken, "w", ZIP_DEFLATED) as result:
        for name in archive.namelist():
            result.writestr(name, b"<broken" if name == "xl/workbook.xml" else archive.read(name))
    response = client.post(
        "/admin/import/preview",
        data={"csrf_token": csrf(client)},
        files={
            "file": (
                "broken.xlsx",
                broken.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 422


def test_xlsx_never_silently_truncates_late_rows():
    import pytest

    from app.services.imports import parse_price_file

    book = Workbook()
    book.active.append(["sku", "name", "category", "price"])
    book.active.append(["GOOD", "Труба", "Трубы", 100])
    book.active.cell(1003, 1, "BAD")
    book.active.cell(1003, 2, "Неверная цена")
    book.active.cell(1003, 3, "Трубы")
    book.active.cell(1003, 4, -1)
    output = BytesIO()
    book.save(output)
    with pytest.raises(ValueError):
        parse_price_file("price.xlsx", output.getvalue())


def test_import_uses_first_worksheet():
    from app.services.imports import parse_price_file

    book = Workbook()
    for sheet, sku in [(book.active, "FIRST"), (book.create_sheet("Second"), "SECOND")]:
        sheet.append(["sku", "name", "category", "price"])
        sheet.append([sku, "Труба", "Трубы", 100])
    book.active = 1
    output = BytesIO()
    book.save(output)
    assert parse_price_file("price.xlsx", output.getvalue())[0]["sku"] == "FIRST"


def test_oversized_csv_cell_returns_validation_error(store):
    client, factory = store
    login_admin(client, factory)
    content = ("sku,name,category,price,description\nBIG,Труба,Трубы,100," + "x" * 140000).encode()
    response = client.post(
        "/admin/import/preview",
        data={"csrf_token": csrf(client)},
        files={"file": ("price.csv", content, "text/csv")},
    )
    assert response.status_code == 422
